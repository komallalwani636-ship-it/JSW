"""Generator for comprehensive CPL-2 scheduling test dataset.

Creates an Excel spreadsheet (.xlsx) with sheet 'Hrstock Report' matching
the exact 84-column format of standard JSW production reports, containing
every pattern, exclusion rule, and violation condition for complete system testing.
"""

from __future__ import annotations

import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADERS_84 = [
    "HR Order No", "HR Coil No", "Thk", "Wdt", "Wgt", "Specific Weight", "Grade",
    "HR Received Date", "CPLAgeDays", "CPLAgeHours", "Days", "Age Hours", "Status",
    "Cust", "Orders", "CR Coil No", "Ord Thk", "Ord Wdt", "Ord Crgrade", "Tdc",
    "Unit", "Loading Port", "Discharging Port", "Edge Condth", "Ord Thk Min",
    "Ord Thk Max", "Ord Wdth Min", "Ord Wdth Max", "Routing", "Product",
    "Coil Saddle Loc", "SPCL Customer", "Work Order No", "Purchase Coil No",
    "Supplier Name", "Oiling Code", "Ord Ra Value", "HR Production Date", "HSM Unit",
    "Material", "NCO FLag", "EQ Spec", "End App", "PPC Remarks", "Insp Remarks",
    "Receiving Remarks", "Doc Type", "Heat No", "Sample Coil", "Coating Type",
    "Coating Code", "Segment", "PrimeSlit Doc Type", "Act Path", "Prev Unit",
    "FT", "CT", "YS", "UTS", "EL", "Ref custname", "Taget Width", "Silicon %",
    "Promise Date", "Schedule Ship Date", "Next Work Center", "Order_Min_Weight",
    "Order_Max_Weight", "Schd_line_no", "NRI", "FP_FIRST_COMMIT_DATE",
    "EARLIEST_START_DATETIME", "LATEST_START_DATETIME", "PLANNEDSTARTDATETIME",
    "PLANNEDENDDATETIME", "QTYPRODUCED", "WIDTH", "CAMPAIGN", "SO_Creation_date",
    "SO_Age", "Order_status", "S_NCO", "S_UD_CODE", "S_Q_LEVEL",
]


def make_row(
    coil_no: str,
    thk: float | None = 2.0,
    wdt: float | None = 1200.0,
    wgt: float | None = 20.0,
    grade: str = "JVHTR01AE0",
    tdc: str = "JVPTR01AE0",
    product: str = "HRPO",
    age_hours: float | None = 100.0,
    status: str = "TA",
    order_status: str = "O",
    act_path: any = 0,
    prev_unit: any = 0,
    silicon_pct: float | None = 0.05,
    next_work_center: str | None = None,
    order_no: str = "0403001001-10",
    receiving_remarks: str = "TEST COIL",
) -> list:
    """Create an 84-column row matching standard report format."""
    row = [None] * len(HEADERS_84)
    row[0] = order_no
    row[1] = coil_no
    row[2] = thk
    row[3] = wdt
    row[4] = wgt
    row[5] = round((wgt or 0) / ((wdt or 1) / 1000.0), 3) if wgt and wdt else 16.0
    row[6] = grade
    row[7] = "2026-09-01"
    row[11] = age_hours
    row[12] = status
    row[13] = "TEST_CUSTOMER"
    row[15] = f"CR_{coil_no}"
    row[19] = tdc
    row[28] = "CPL2"
    row[29] = product
    row[45] = receiving_remarks
    row[53] = act_path
    row[54] = prev_unit
    row[61] = wdt
    row[62] = silicon_pct
    row[65] = next_work_center
    row[80] = order_status
    return row


def build_test_dataset(output_path: str = "TEST_CPL2_COMPREHENSIVE_DATASET.xlsx") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hrstock Report"

    # Row 1: Title
    title_row = ["DCN NO CRM-PPC-QR-005  Format No CRM-PPC-FR-005"] + [None] * (len(HEADERS_84) - 1)
    ws.append(title_row)

    # Row 2: Headers
    ws.append(HEADERS_84)

    rows: list[list] = []
    coil_counter = 1000

    def next_id(prefix="TC"):
        nonlocal coil_counter
        coil_counter += 1
        return f"{prefix}{coil_counter}"

    # =========================================================================
    # SECTION 1: Standard Eligible HRPO Coils (Smooth Profile)
    # =========================================================================
    for thk in [1.6, 1.8, 2.0, 2.5, 3.0, 3.5, 4.0]:
        for wdt in [1200.0, 1220.0, 1240.0, 1250.0]:
            rows.append(make_row(
                coil_no=next_id("HRPO_OK"),
                thk=thk,
                wdt=wdt,
                wgt=21.5,
                product="HRPO",
                age_hours=90.0,
                grade="JVHTR01AE0",
                tdc="JVPTR01AE0",
                receiving_remarks="Eligible standard HRPO",
            ))

    # =========================================================================
    # SECTION 2: Eligible HRSPO Coils (Various Age Milestones)
    # =========================================================================
    # Age > 72h
    for i in range(4):
        rows.append(make_row(
            coil_no=next_id("HRSPO_72"),
            thk=2.2,
            wdt=1150.0 + (i * 20),
            product="HRSPO",
            age_hours=85.0,
            receiving_remarks="HRSPO Age > 72h",
        ))
    # Age > 120h
    for i in range(4):
        rows.append(make_row(
            coil_no=next_id("HRSPO_120"),
            thk=2.2,
            wdt=1160.0 + (i * 15),
            product="HRSPO",
            age_hours=145.0,
            receiving_remarks="HRSPO Age > 120h",
        ))
    # Age > 168h (Critical 7-day FIFO priority)
    for i in range(4):
        rows.append(make_row(
            coil_no=next_id("HRSPO_168"),
            thk=2.2,
            wdt=1170.0 + (i * 10),
            product="HRSPO",
            age_hours=260.0,
            receiving_remarks="HRSPO Critical Age > 168h",
        ))

    # =========================================================================
    # SECTION 3: Eligible NGO FP Coils (Silicon Matrix)
    # =========================================================================
    # Ultra Low Silicon (< 0.1%)
    for i in range(3):
        rows.append(make_row(
            coil_no=next_id("NGO_UL"),
            thk=1.8,
            wdt=1050.0,
            product="NGO FP",
            grade="JNH250A00",
            silicon_pct=0.04,
            age_hours=110.0,
            receiving_remarks="NGO Ultra-Low Silicon",
        ))
    # Low Silicon (0.1% - 0.5%)
    for i in range(3):
        rows.append(make_row(
            coil_no=next_id("NGO_LO"),
            thk=1.8,
            wdt=1060.0,
            product="NGO FP",
            grade="JNH300A00",
            silicon_pct=0.25,
            age_hours=115.0,
            receiving_remarks="NGO Low Silicon",
        ))
    # Medium Silicon (0.5% - 1.5%)
    for i in range(3):
        rows.append(make_row(
            coil_no=next_id("NGO_MED"),
            thk=1.8,
            wdt=1070.0,
            product="NGO FP",
            grade="JNH400A00",
            silicon_pct=0.85,
            age_hours=120.0,
            receiving_remarks="NGO Medium Silicon",
        ))
    # High Silicon (> 1.5%)
    for i in range(3):
        rows.append(make_row(
            coil_no=next_id("NGO_HI"),
            thk=1.8,
            wdt=1080.0,
            product="NGO FP",
            grade="JNH600A00",
            silicon_pct=2.10,
            age_hours=125.0,
            receiving_remarks="NGO High Silicon",
        ))

    # =========================================================================
    # SECTION 4: Critical TDC Grades (Tests TDC Adjacency Constraint)
    # =========================================================================
    tdc_grades = ["JVPFB60AJS", "JVPTR14AJS", "JVPST01C00", "JVPTR15AJS", "JVPTR13AJS"]
    for i, tdc_code in enumerate(tdc_grades * 2):
        rows.append(make_row(
            coil_no=next_id("TDC_CRIT"),
            thk=4.5 + (i * 0.1),
            wdt=1100.0,
            product="HRPO",
            grade=tdc_code.replace("P", "H"),
            tdc=tdc_code,
            age_hours=130.0,
            receiving_remarks=f"Critical TDC grade {tdc_code}",
        ))

    # =========================================================================
    # =========================================================================
    # SECTION 5: APL7 Routing (Tracked coils)
    # =========================================================================
    for i in range(6):
        rows.append(make_row(
            coil_no=next_id("APL7_ROUTED"),
            thk=2.5,
            wdt=1250.0,
            product="HRPO",
            next_work_center="APL7",
            receiving_remarks="APL7",
        ))

    # =========================================================================
    # SECTION 6: Explicit Violations Injection (For testing detector & optimizer)
    # =========================================================================
    # V1: Thin Width Jump >= 150mm (e.g. 1260mm to 1020mm = 240mm jump)
    rows.append(make_row(
        coil_no=next_id("VIOL_WDT_THIN_A"),
        thk=1.6,
        wdt=1260.0,
        product="HRPO",
        receiving_remarks="Thin Width Jump Start (1260mm)",
    ))
    rows.append(make_row(
        coil_no=next_id("VIOL_WDT_THIN_B"),
        thk=1.6,
        wdt=1020.0,
        product="HRPO",
        receiving_remarks="Thin Width Jump (1020mm, diff 240mm >= 150mm)",
    ))

    # V2: Thick Width Jump >= 300mm (e.g. 1550mm to 1100mm = 450mm jump)
    rows.append(make_row(
        coil_no=next_id("VIOL_WDT_THICK_A"),
        thk=3.5,
        wdt=1550.0,
        product="HRPO",
        receiving_remarks="Thick Width Jump Start (1550mm)",
    ))
    rows.append(make_row(
        coil_no=next_id("VIOL_WDT_THICK_B"),
        thk=3.5,
        wdt=1100.0,
        product="HRPO",
        receiving_remarks="Thick Width Jump (1100mm, diff 450mm >= 300mm)",
    ))

    # V3: Thin Thickness Step > 0.4mm (1.5mm to 1.95mm = 0.45mm diff)
    rows.append(make_row(
        coil_no=next_id("VIOL_THK_THIN_A"),
        thk=1.5,
        wdt=1200.0,
        product="HRPO",
        receiving_remarks="Thin Thickness Step (1.5mm)",
    ))
    rows.append(make_row(
        coil_no=next_id("VIOL_THK_THIN_B"),
        thk=1.95,
        wdt=1200.0,
        product="HRPO",
        receiving_remarks="Thin Thickness Step (1.95mm, diff 0.45mm > 0.4mm)",
    ))

    # V4: Thick Thickness Step > 1.0mm (2.5mm to 4.2mm = 1.7mm diff)
    rows.append(make_row(
        coil_no=next_id("VIOL_THK_THICK_A"),
        thk=2.5,
        wdt=1200.0,
        product="HRPO",
        receiving_remarks="Thick Thickness Step (2.5mm)",
    ))
    rows.append(make_row(
        coil_no=next_id("VIOL_THK_THICK_B"),
        thk=4.2,
        wdt=1200.0,
        product="HRPO",
        receiving_remarks="Thick Thickness Step (4.2mm, diff 1.7mm > 1.0mm)",
    ))

    # =========================================================================
    # SECTION 7: Excluded Coils - Non-CPL2 Products
    # =========================================================================
    for prod in ["CR", "GP", "GI"]:
        rows.append(make_row(
            coil_no=next_id("EXCL_PROD"),
            product=prod,
            receiving_remarks=f"Excluded: Product {prod}",
        ))

    # =========================================================================
    # SECTION 8: Excluded Coils - Act Path / Prev Unit (Routing Exclusions)
    # =========================================================================
    # String routing codes (produce Parse Warnings + Exclusion)
    for code, unit in [("TTSJ", "SPCL"), ("W", "SPCL"), ("TTJ", "SPCL"), ("T", "CPL2"), ("WT", "CPL2")]:
        rows.append(make_row(
            coil_no=next_id("EXCL_STR_ROUTING"),
            act_path=code,
            prev_unit=unit,
            receiving_remarks=f"Excluded: String Act Path {code} / Prev Unit {unit}",
        ))
    # Numeric non-zero routing codes
    for ap, pu in [(1, 0), (0, 2), (3, 4)]:
        rows.append(make_row(
            coil_no=next_id("EXCL_NUM_ROUTING"),
            act_path=ap,
            prev_unit=pu,
            receiving_remarks=f"Excluded: Non-zero Act Path {ap} / Prev Unit {pu}",
        ))

    # =========================================================================
    # SECTION 9: Excluded Coils - Status Exclusions
    # =========================================================================
    for st in ["HOLD", "REJ", "BLOCKED", "SCRAP"]:
        rows.append(make_row(
            coil_no=next_id("EXCL_STATUS"),
            status=st,
            receiving_remarks=f"Excluded: Non-TA status {st}",
        ))

    # =========================================================================
    # SECTION 10: Excluded Coils - Closed/Hold Orders
    # =========================================================================
    for ost in ["CLOSED", "CLOSED", "CLOSED"]:
        rows.append(make_row(
            coil_no=next_id("EXCL_ORDER"),
            order_status=ost,
            receiving_remarks=f"Excluded: Closed order status {ost}",
        ))


    # =========================================================================
    # SECTION 11: Excluded Coils - Fresh Age Rule (<= 72 Hours)
    # =========================================================================
    for age in [12.0, 24.0, 48.0, 72.0]:
        rows.append(make_row(
            coil_no=next_id("EXCL_AGE"),
            age_hours=age,
            receiving_remarks=f"Excluded: Age <= 72h ({age}h)",
        ))

    # Append all rows to sheet
    for r in rows:
        ws.append(r)

    # Style header row for clean aesthetics
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col in range(1, len(HEADERS_84) + 1):
        cell = ws.cell(row=2, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(output_path)
    print(f"Generated test dataset at: {os.path.abspath(output_path)}")
    print(f"Total rows written: {len(rows)}")


if __name__ == "__main__":
    build_test_dataset()
