"""XLSX exporter for the CPL-2 scheduling system.

Builds a workbook with Plan and APL7 sheets matching the manual export format.

# Feature: cpl2-scheduling-system
# Validates: Requirements 10.2, 10.3, 10.4, 10.5
"""

from __future__ import annotations

from io import BytesIO
from typing import Protocol

from openpyxl import Workbook

from engine.parser import REQUIRED_COLUMNS


class _ExportCoil(Protocol):
    all_columns: dict


def _resolve_headers(coils: list[_ExportCoil]) -> list[str]:
    if not coils:
        return list(REQUIRED_COLUMNS)
    first_keys = list(coils[0].all_columns.keys())
    if first_keys:
        return first_keys
    return list(REQUIRED_COLUMNS)


def _write_coil_sheet(
    ws,
    coils: list[_ExportCoil],
    headers: list[str],
    *,
    is_draft: bool = False,
) -> None:
    header_row = 2 if is_draft else 1
    data_start = header_row + 1

    if is_draft:
        ws["A1"] = "DRAFT"

    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)

    for row_offset, coil in enumerate(coils):
        row_idx = data_start + row_offset
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=coil.all_columns.get(header, ""))


def build_export_workbook(
    eligible_coils: list[_ExportCoil],
    apl7_coils: list[_ExportCoil],
    is_draft: bool,
) -> BytesIO:
    """Build an XLSX workbook with Plan and APL7 sheets."""
    wb = Workbook()
    plan_ws = wb.active
    plan_ws.title = "Plan"

    headers = _resolve_headers(eligible_coils or apl7_coils)
    _write_coil_sheet(plan_ws, eligible_coils, headers, is_draft=is_draft)

    apl7_ws = wb.create_sheet("APL7")
    _write_coil_sheet(apl7_ws, apl7_coils, headers, is_draft=False)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
