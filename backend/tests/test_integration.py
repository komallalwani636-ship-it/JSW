"""Full-stack integration tests using the real HR Stock XLS file."""

from __future__ import annotations

import io
import os
import sys

import openpyxl
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))

from conftest import TestingSessionLocal, init_test_db, override_get_db

init_test_db()

from api.auth_service import create_access_token, hash_password
from db.models import User

XLS_FILE = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "CPL2 HR STOCK FOR SCHEDULING.xls")
)


def _build_client() -> TestClient:
    from api.dependencies import get_db
    from main import app

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app, raise_server_exceptions=False)


def _auth_headers() -> dict[str, str]:
    db = TestingSessionLocal()
    try:
        user = db.query(User).filter(User.username == "integration_planner").one_or_none()
        if user is None:
            user = User(
                username="integration_planner",
                password_hash=hash_password("password"),
                role="planner",
            )
            db.add(user)
            db.commit()
        token = create_access_token({"sub": "integration_planner", "role": "planner"})
        return {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


@pytest.mark.skipif(not os.path.exists(XLS_FILE), reason="Real XLS file not available")
def test_full_api_flow_with_real_xls() -> None:
    client = _build_client()
    headers = _auth_headers()

    with open(XLS_FILE, "rb") as fh:
        upload = client.post(
            "/api/uploads",
            files={"file": ("CPL2 HR STOCK FOR SCHEDULING.xls", fh, "application/vnd.ms-excel")},
            headers=headers,
        )
    assert upload.status_code == 200, upload.text
    upload_body = upload.json()
    upload_id = upload_body["upload_id"]
    summary = upload_body["parse_summary"]
    assert summary["total_rows"] > 0
    assert summary["eligible_count"] > 0

    generate = client.post("/api/schedules", json={"upload_id": upload_id}, headers=headers)
    assert generate.status_code == 200, generate.text
    schedule_id = generate.json()["id"]
    assert generate.json()["status"] == "draft"

    kpis = client.get(f"/api/schedules/{schedule_id}/kpis", headers=headers)
    assert kpis.status_code == 200
    assert kpis.json()["total_eligible"] > 0

    items = client.get(f"/api/schedules/{schedule_id}/items", headers=headers).json()
    eligible_items = [item for item in items if not item["is_apl7"]]
    assert len(eligible_items) >= 2
    coil_ids = [item["coil"]["id"] for item in eligible_items]
    reordered = [coil_ids[1], coil_ids[0]] + coil_ids[2:]

    reorder = client.patch(
        f"/api/schedules/{schedule_id}/reorder",
        json={"coil_ids": reordered},
        headers=headers,
    )
    assert reorder.status_code == 200
    assert "violations" in reorder.json()

    violations = reorder.json().get("violations", [])
    finalize_body = (
        {"override_reason": "Integration test approval"}
        if violations
        else {}
    )
    finalize = client.post(
        f"/api/schedules/{schedule_id}/finalize",
        json=finalize_body,
        headers=headers,
    )
    assert finalize.status_code == 200
    assert finalize.json()["status"] == "final"

    locked = client.patch(
        f"/api/schedules/{schedule_id}/reorder",
        json={"coil_ids": reordered},
        headers=headers,
    )
    assert locked.status_code == 409

    unlock = client.post(f"/api/schedules/{schedule_id}/unlock", headers=headers)
    assert unlock.status_code == 200
    assert unlock.json()["status"] == "draft"

    export = client.get(f"/api/schedules/{schedule_id}/export", headers=headers)
    assert export.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(export.content))
    assert "Plan" in wb.sheetnames
    assert "APL7" in wb.sheetnames
    assert wb["Plan"].max_row >= 2
    assert wb["APL7"].max_row >= 1
