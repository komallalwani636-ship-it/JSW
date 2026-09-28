"""Property-based tests for schedule export.

# Feature: cpl2-scheduling-system, Property 25: Export contains all eligible coils
"""

from __future__ import annotations

import os
import sys
from io import BytesIO

import openpyxl
from hypothesis import given, settings, HealthCheck
from hypothesis import strategies as st

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from engine.exporter import build_export_workbook
from engine.parser import REQUIRED_COLUMNS


class _FakeCoil:
    def __init__(self, hr_coil_no: str):
        self.all_columns = {col: "" for col in REQUIRED_COLUMNS}
        self.all_columns["HR Coil No"] = hr_coil_no


@given(
    eligible=st.lists(
        st.text(min_size=1, max_size=8, alphabet="ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"),
        min_size=0,
        max_size=5,
        unique=True,
    ),
    apl7=st.lists(
        st.text(min_size=1, max_size=8, alphabet="ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"),
        min_size=0,
        max_size=3,
        unique=True,
    ),
)
@settings(max_examples=3, deadline=None, suppress_health_check=list(HealthCheck), database=None)
def test_export_contains_all_eligible_coils(eligible: list[str], apl7: list[str]) -> None:
    eligible_set = set(eligible)
    apl7_set = set(apl7) - eligible_set
    eligible = list(eligible_set)
    apl7 = list(apl7_set)

    eligible_coils = [_FakeCoil(no) for no in eligible]
    apl7_coils = [_FakeCoil(no) for no in apl7]

    workbook_bytes = build_export_workbook(eligible_coils, apl7_coils, is_draft=False)
    wb = openpyxl.load_workbook(BytesIO(workbook_bytes.getvalue()))

    assert "Plan" in wb.sheetnames
    assert "APL7" in wb.sheetnames

    plan_ws = wb["Plan"]
    plan_headers = [cell.value for cell in plan_ws[1]]
    hr_col_idx = plan_headers.index("HR Coil No") + 1
    exported_eligible = {
        row[hr_col_idx - 1].value
        for row in plan_ws.iter_rows(min_row=2, values_only=False)
        if row[hr_col_idx - 1].value
    }
    assert exported_eligible == set(eligible)

    apl7_ws = wb["APL7"]
    apl7_headers = [cell.value for cell in apl7_ws[1]]
    apl7_col_idx = apl7_headers.index("HR Coil No") + 1
    exported_apl7 = {
        row[apl7_col_idx - 1].value
        for row in apl7_ws.iter_rows(min_row=2, values_only=False)
        if row[apl7_col_idx - 1].value
    }
    assert exported_apl7 == set(apl7)
